"""Text extraction from attachments so their contents are searchable."""

import io
import logging
import signal
import sys
import threading
import warnings
import zipfile
from contextlib import contextmanager
from email import policy
from email.parser import BytesParser

from .textutil import html_to_text, tidy_text

EXTRACTOR_VERSION = 1  # bump when extraction changes, to invalidate the cached attachment text
MAX_CHARS = 200_000
TIME_LIMIT = 60  # seconds per attachment; malformed PDFs can make parsers spin
MAX_PDF_PAGES = 300
MAX_SHEET_ROWS = 5_000

logging.getLogger("pypdf").setLevel(logging.ERROR)

TEXT_EXT = {"txt", "csv", "tsv", "md", "log", "json", "xml", "ics", "vcf", "yaml", "yml"}
EXTRACTABLE_EXT = TEXT_EXT | {"pdf", "docx", "xlsx", "xlsm", "pptx", "htm", "html", "eml", "rtf"}


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        reader = PdfReader(io.BytesIO(data), strict=False)
        out = []
        for i, page in enumerate(reader.pages):
            if i >= MAX_PDF_PAGES:
                break
            out.append(page.extract_text() or "")
            if sum(map(len, out)) > MAX_CHARS:
                break
    return "\n\n".join(out)


def _docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append(" | ".join(c.text for c in row.cells))
    return "\n".join(out)


def _xlsx(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"# {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_SHEET_ROWS:
                break
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    out = []
    for n, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
        out.append(f"# Slide {n}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
    return "\n".join(out)


def _rtf(data: bytes) -> str:
    import re

    s = data.decode("latin-1", "replace")
    s = re.sub(r"\\'([0-9a-f]{2})", lambda m: bytes.fromhex(m.group(1)).decode("cp1252", "replace"), s)
    s = re.sub(r"\\par[d]?\b|\\line\b", "\n", s)
    s = re.sub(r"\{\\\*[^{}]*\}|\\[a-z]+-?\d* ?|[{}]", "", s)
    return s


def _eml(data: bytes) -> str:
    from .parse import header  # local import: parse imports textutil only

    msg = BytesParser(policy=policy.default).parsebytes(data)
    head = "\n".join(f"{h}: {header(msg, h)}" for h in ("From", "To", "Cc", "Date", "Subject") if header(msg, h))
    body = ""
    try:
        part = msg.get_body(("plain", "html"))
        if part is not None:
            content = part.get_content()
            body = html_to_text(content) if part.get_content_type() == "text/html" else content
    except Exception:
        pass
    return head + "\n\n" + body


@contextmanager
def _time_limit(seconds: int):
    if not hasattr(signal, "SIGALRM") or threading.current_thread() is not threading.main_thread():
        yield
        return

    def stop(signum, frame):
        raise TimeoutError

    old = signal.signal(signal.SIGALRM, stop)
    signal.setitimer(signal.ITIMER_REAL, seconds)
    try:
        yield
    finally:
        signal.setitimer(signal.ITIMER_REAL, 0)
        signal.signal(signal.SIGALRM, old)


def extract_text(name: str, mime: str, data: bytes) -> str:
    try:
        with _time_limit(TIME_LIMIT):
            return _extract(name, mime, data)
    except TimeoutError:
        print(f"  ! gave up extracting text from {name} after {TIME_LIMIT}s", file=sys.stderr)
        return ""


def _extract(name: str, mime: str, data: bytes) -> str:
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    try:
        if ext == "pdf" or mime == "application/pdf":
            text = _pdf(data)
        elif ext == "docx":
            text = _docx(data)
        elif ext in ("xlsx", "xlsm"):
            text = _xlsx(data)
        elif ext == "pptx":
            text = _pptx(data)
        elif ext in ("htm", "html") or mime == "text/html":
            text = html_to_text(data.decode("utf-8", "replace"))
        elif ext == "eml" or mime == "message/rfc822":
            text = _eml(data)
        elif ext == "rtf" or mime in ("application/rtf", "text/rtf"):
            text = _rtf(data)
        elif ext in TEXT_EXT or mime.startswith("text/"):
            text = data.decode("utf-8", "replace")
        else:
            return ""
    except TimeoutError:
        raise
    except (Exception, zipfile.BadZipFile):
        return ""
    return tidy_text(text)[:MAX_CHARS]
